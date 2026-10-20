import logging

from rest_framework import status, viewsets, permissions
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import action
from django.shortcuts import get_object_or_404
from django.db import close_old_connections
from django.http import HttpResponse
from io import BytesIO
from threading import Thread
from rag_app.models import DocumentCollection, Document, Conversation, Message, Bookmark, Feedback, EvaluationResult
from rag_app.services.rag_service import RAGService
from rag_app.services.analysis_service import AnalysisService
from rag_app.services.comparison_service import ComparisonService
from rag_app.services.search_service import SearchService
from .serializers import (
    DocumentCollectionSerializer, DocumentSerializer, DocumentChunkSerializer,
    ConversationSerializer, MessageSerializer, BookmarkSerializer, FeedbackSerializer,
    DocumentUploadSerializer, ChatQuerySerializer, SummaryRequestSerializer,
    SearchRequestSerializer, CompareRequestSerializer, EvaluationResultSerializer
)

logger = logging.getLogger(__name__)


def _process_document_async(document_id):
    close_old_connections()
    try:
        RAGService().process_document(document_id)
    finally:
        close_old_connections()

class HealthCheckView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        return Response({
            "status": "healthy",
            "service": "AI Document Intelligence & Enterprise RAG Assistant",
            "version": "2.0.0"
        })


class ExportAPIView(APIView):
    """Export an owned answer or complete conversation as a private download."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        export_format = request.data.get('format', 'pdf')
        if export_format not in {'pdf', 'docx'}:
            return Response({'error': 'format must be pdf or docx'}, status=status.HTTP_400_BAD_REQUEST)
        message_id = request.data.get('message_id')
        conversation_id = request.data.get('conversation_id')
        if bool(message_id) == bool(conversation_id):
            return Response({'error': 'Provide exactly one message_id or conversation_id.'}, status=status.HTTP_400_BAD_REQUEST)
        if message_id:
            messages = Message.objects.filter(pk=message_id, conversation__user=request.user)
        else:
            messages = Message.objects.filter(conversation_id=conversation_id, conversation__user=request.user).order_by('created_at')
        rows = list(messages.values_list('role', 'content', 'sources'))
        if not rows:
            return Response({'error': 'No owned answer or conversation was found.'}, status=status.HTTP_404_NOT_FOUND)
        content = '\n\n'.join(f'{role.title()}: {text}' for role, text, _ in rows)
        output = BytesIO()
        if export_format == 'docx':
            from docx import Document as WordDocument
            document = WordDocument()
            document.add_heading('Document Intelligence Conversation', 0)
            for role, text, sources in rows:
                document.add_heading(role.title(), level=2)
                document.add_paragraph(text)
                for source in sources or []:
                    document.add_paragraph(
                        f"Source: {source.get('document_title', 'Document')}, page {source.get('page_number', 'N/A')} — "
                        f"{source.get('passage', source.get('chunk_preview', ''))}"
                    )
            document.save(output)
            mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            import fitz
            import textwrap
            pdf = fitz.open()
            page = pdf.new_page()
            y = 50
            for line in content.splitlines():
                for wrapped in textwrap.wrap(line, width=95) or ['']:
                    if y > 770:
                        page = pdf.new_page()
                        y = 50
                    page.insert_text((50, y), wrapped, fontsize=10)
                    y += 15
            for role, _, sources in rows:
                if role == 'assistant' and sources:
                    for source in sources:
                        if y > 740:
                            page = pdf.new_page()
                            y = 50
                        citation = f"Source: {source.get('document_title', 'Document')}, page {source.get('page_number', 'N/A')}"
                        page.insert_textbox(fitz.Rect(50, y, 545, min(y + 50, 790)), citation, fontsize=9)
                        y += 30
            pdf.save(output)
            pdf.close()
            mime = 'application/pdf'
        output.seek(0)
        response = HttpResponse(output.read(), content_type=mime)
        response['Content-Disposition'] = f'attachment; filename="conversation.{export_format}"'
        response['X-Content-Type-Options'] = 'nosniff'
        return response


class DocumentCollectionViewSet(viewsets.ModelViewSet):
    serializer_class = DocumentCollectionSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return DocumentCollection.objects.filter(user=self.request.user)

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)


class DocumentViewSet(viewsets.ModelViewSet):
    serializer_class = DocumentSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return Document.objects.filter(user=self.request.user)

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)

    def perform_destroy(self, instance):
        file_field = instance.file
        instance.delete()  # cascades chunks first, preserving tenant ownership checks
        if file_field:
            file_field.delete(save=False)
        from rag_app.services.vector_service import VectorService
        VectorService.rebuild_from_database()

    @action(detail=True, methods=['post'])
    def reprocess(self, request, pk=None):
        doc = self.get_object()
        doc.status = 'uploaded'
        doc.processing_error = ''
        doc.save(update_fields=['status', 'processing_error'])
        Thread(target=_process_document_async, args=(doc.id,), daemon=True).start()
        return Response({
            "success": True,
            "document": DocumentSerializer(doc).data,
            "message": "Document reprocessing started."
        })

    @action(detail=True, methods=['get'])
    def analyze(self, request, pk=None):
        doc = self.get_object()
        result = AnalysisService.analyze_document_and_generate_questions(doc.id)
        return Response(result)


class DocumentUploadAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = DocumentUploadSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        uploaded_file = serializer.validated_data['file']
        collection_id = serializer.validated_data.get('collection_id')
        collection = None
        if collection_id:
            collection = get_object_or_404(DocumentCollection, id=collection_id, user=request.user)

        title = serializer.validated_data.get('title') or uploaded_file.name

        doc = Document.objects.create(
            user=request.user,
            collection=collection,
            title=title,
            file=uploaded_file,
            file_size=uploaded_file.size,
            status='uploaded'
        )

        Thread(target=_process_document_async, args=(doc.id,), daemon=True).start()

        return Response({
            "success": True,
            "document": DocumentSerializer(doc).data,
            "suggested_questions": doc.suggested_questions,
            "message": "Document uploaded; ingestion started. Poll the document endpoint for status."
        }, status=status.HTTP_202_ACCEPTED)


class ChatAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = ChatQuerySerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        question = serializer.validated_data['question']
        conversation_id = serializer.validated_data.get('conversation_id')
        document_id = serializer.validated_data.get('document_id')
        collection_id = serializer.validated_data.get('collection_id')

        if document_id and not Document.objects.filter(
            id=document_id, user=request.user, status='completed'
        ).exists():
            return Response({'error': 'The selected document is not available.'}, status=status.HTTP_404_NOT_FOUND)
        if collection_id and not DocumentCollection.objects.filter(
            id=collection_id, user=request.user
        ).exists():
            return Response({'error': 'The selected collection is not available.'}, status=status.HTTP_404_NOT_FOUND)
        if conversation_id and not Conversation.objects.filter(
            id=conversation_id, user=request.user
        ).exists():
            return Response({'error': 'The selected conversation is not available.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            rag_service = RAGService()
            result = rag_service.query_rag(
                question=question,
                user_id=request.user.id,
                conversation_id=conversation_id,
                document_id=document_id,
                collection_id=collection_id
            )
        except Exception:
            logger.exception('Chat request failed for user_id=%s', request.user.id)
            return Response(
                {'error': 'The chat service is temporarily unavailable. Please try again.'},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        return Response(result, status=status.HTTP_200_OK)


class SearchAPIView(APIView):
    """Executes standalone Hybrid Search (FAISS + BM25) for inspection."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = SearchRequestSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        values = serializer.validated_data

        result = SearchService.execute_search(
            query=values['query'],
            user_id=request.user.id,
            document_id=values.get('document_id'),
            collection_id=values.get('collection_id')
        )

        return Response(result, status=status.HTTP_200_OK)


class ConversationViewSet(viewsets.ModelViewSet):
    serializer_class = ConversationSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return Conversation.objects.filter(user=self.request.user)


class SummaryAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = SummaryRequestSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        document_id = serializer.validated_data['document_id']
        mode = serializer.validated_data['mode']

        rag_service = RAGService()
        summary_result = rag_service.generate_summary(document_id, request.user.id, mode=mode)
        summary_status = status.HTTP_200_OK if summary_result["success"] else status.HTTP_400_BAD_REQUEST

        return Response(summary_result, status=summary_status)


class CompareAPIView(APIView):
    """Executes multi-document comparison analysis between Document A and Document B."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = CompareRequestSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        values = serializer.validated_data

        result = ComparisonService.compare_documents(values['doc_a_id'], values['doc_b_id'], request.user.id)
        res_status = status.HTTP_200_OK if result["success"] else status.HTTP_400_BAD_REQUEST

        return Response(result, status=res_status)


class BookmarkViewSet(viewsets.ModelViewSet):
    """CRUD API for saving and managing user bookmarks."""
    serializer_class = BookmarkSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return Bookmark.objects.filter(user=self.request.user)

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)


class FeedbackAPIView(APIView):
    """Submits user rating (+1 / -1) and feedback comments on assistant answers."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        message_id = request.data.get('message_id')
        rating = request.data.get('rating')
        comment = request.data.get('comment', '')

        if not message_id or rating not in [1, -1]:
            return Response({"error": "Valid message_id and rating (+1 or -1) are required."}, status=status.HTTP_400_BAD_REQUEST)

        msg = get_object_or_404(Message, id=message_id, conversation__user=request.user)

        feedback, created = Feedback.objects.update_or_create(
            user=request.user,
            message=msg,
            defaults={'rating': rating, 'comment': comment}
        )

        return Response({
            "success": True,
            "feedback_id": feedback.id,
            "rating": feedback.rating,
            "message": "Feedback recorded successfully."
        }, status=status.HTTP_200_OK)


class EvaluationResultViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = EvaluationResultSerializer
    permission_classes = [permissions.IsAuthenticated]
    queryset = EvaluationResult.objects.all()
