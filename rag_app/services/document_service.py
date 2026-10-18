"""Document text extraction adapters. Optional parsers are loaded only for their formats."""
import csv
from pathlib import Path


class DocumentExtractionService:
    @staticmethod
    def extract(path):
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix == '.pdf':
            from .pdf_service import PDFService
            result = PDFService.extract_text_from_pdf(str(file_path))
            if not result['success']:
                raise ValueError(result['error'])
            pages = result['pages']
            if not any(p['text'].strip() for p in pages):
                pages = DocumentExtractionService._ocr_pdf(file_path)
            return {'pages': pages}
        if suffix in {'.txt', '.md'}:
            text = file_path.read_text(encoding='utf-8-sig', errors='replace')
            return {'pages': [{'page': 1, 'text': text}]}
        if suffix == '.csv':
            with file_path.open(encoding='utf-8-sig', newline='') as source:
                text = '\n'.join(' | '.join(row) for row in csv.reader(source))
            return {'pages': [{'page': 1, 'text': text}]}
        if suffix == '.docx':
            from docx import Document as WordDocument
            doc = WordDocument(str(file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            paragraphs += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
            return {'pages': [{'page': 1, 'text': '\n\n'.join(paragraphs)}]}
        if suffix == '.pptx':
            from pptx import Presentation
            presentation = Presentation(str(file_path))
            pages = []
            for number, slide in enumerate(presentation.slides, 1):
                texts = [shape.text for shape in slide.shapes if getattr(shape, 'has_text_frame', False) and shape.text.strip()]
                texts += [' | '.join(cell.text for row in shape.table.rows for cell in row.cells)
                          for shape in slide.shapes if getattr(shape, 'has_table', False)]
                pages.append({'page': number, 'text': '\n'.join(texts)})
            return {'pages': pages}
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(str(file_path), read_only=True, data_only=True)
            pages = []
            for number, sheet in enumerate(workbook.worksheets, 1):
                rows = [' | '.join('' if value is None else str(value) for value in row)
                        for row in sheet.iter_rows(values_only=True)]
                pages.append({'page': number, 'text': f'{sheet.title}\n' + '\n'.join(rows)})
            return {'pages': pages}
        raise ValueError(f'Unsupported document format: {suffix}')

    @staticmethod
    def _ocr_pdf(path):
        try:
            import fitz
            import pytesseract
            from PIL import Image
            from io import BytesIO
        except ImportError as exc:
            raise ValueError('Scanned PDF needs optional OCR dependencies (pytesseract and Pillow).') from exc
        pages = []
        with fitz.open(path) as pdf:
            for number, page in enumerate(pdf, 1):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(BytesIO(pixmap.tobytes('png')))
                pages.append({'page': number, 'text': pytesseract.image_to_string(image)})
        return pages
